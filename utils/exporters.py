import pandas as pd
import json
from docx import Document
from fpdf import FPDF

class Exporter:
    def __init__(self, data):
        self.df = pd.DataFrame(data)

    def to_csv(self, path):
        self.df.to_csv(path, index=False)

    def to_excel(self, path):
        self.df.to_excel(path, index=False, engine='openpyxl')

    def to_json(self, path):
        self.df.to_json(path, orient='records', indent=4)
        
    def to_pdf(self, path):
        pdf = FPDF()
        pdf.set_auto_page_break(auto=True, margin=10)
        pdf.add_page()
        pdf.set_font("Arial", size=9)
        pdf.cell(0, 8, "Weapons Export Report", ln=True)
        pdf.ln(2)
        cols = list(self.df.columns)[:6]
        for _, row in self.df.iterrows():
            parts = [f"{c}: {row.get(c, '')}" for c in cols]
            pdf.multi_cell(0, 6, " | ".join(parts))
            pdf.ln(1)
        pdf.output(path)

    def to_docx(self, path):
        doc = Document()
        doc.add_heading("Weapons Export Report", level=1)
        table = doc.add_table(rows=1, cols=len(self.df.columns))
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        for idx, col in enumerate(self.df.columns):
            hdr[idx].text = str(col)
        for _, row in self.df.iterrows():
            cells = table.add_row().cells
            for idx, col in enumerate(self.df.columns):
                cells[idx].text = str(row.get(col, ""))
        doc.save(path)
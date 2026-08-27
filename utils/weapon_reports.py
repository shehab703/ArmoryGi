"""
Rich weapon reports (DOCX / PDF): individual profiles with images,
range/effects map diagrams, multi-weapon comparisons, category-grouped
lists, and full-specification catalogs.
"""
from __future__ import annotations

import html
import math
import os
import tempfile
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

from docx import Document
from docx.shared import Inches, Pt

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Image as RLImage,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from utils.spec_fields import SPEC_FIELDS, SPEC_FIELDS_AR

# ---------------------------------------------------------------------------
# Geo / visuals (aligned with map view inner ring heuristic)
# ---------------------------------------------------------------------------


def range_ring_coords(
    origin_lat: float,
    origin_lon: float,
    range_km: float,
    num_points: int = 72,
) -> List[Tuple[float, float]]:
    if not range_km or not origin_lat or not origin_lon:
        return []
    earth_radius_km = 6371.0
    coords: List[Tuple[float, float]] = []
    for i in range(num_points):
        angle = 2 * math.pi * i / num_points
        lat_offset = (range_km / earth_radius_km) * math.cos(angle) * (180 / math.pi)
        lon_offset = (
            (range_km / earth_radius_km)
            * math.sin(angle)
            * (180 / math.pi)
            / max(math.cos(math.radians(origin_lat)), 0.2)
        )
        coords.append(
            (
                round(origin_lat + lat_offset, 6),
                round(origin_lon + lon_offset, 6),
            )
        )
    return coords


def build_range_effects_png(weapon: Dict[str, Any], out_path: str, size: Tuple[int, int] = (920, 620)) -> bool:
    """Render origin + min/max range rings to a PNG (Pillow)."""
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        return False

    w, h = size
    im = Image.new("RGB", (w, h), (14, 22, 38))
    draw = ImageDraw.Draw(im)
    ola = float(weapon.get("origin_lat") or 31.5)
    olo = float(weapon.get("origin_lon") or 34.8)
    rmax = float(weapon.get("range_km") or 0)
    rmin = max(rmax * 0.35, 0.0) if rmax else 0.0

    margin = 48
    km_per_deg_lat = 111.0
    km_per_deg_lon = 111.0 * max(math.cos(math.radians(ola)), 0.2)
    extent_km = (rmax * 1.2 + 8.0) if rmax else 45.0
    scale_x = (w - 2 * margin) / (2 * extent_km / km_per_deg_lon)
    scale_y = (h - 2 * margin) / (2 * extent_km / km_per_deg_lat)
    scale = min(scale_x, scale_y)

    def geo_to_px(lat: float, lon: float) -> Tuple[float, float]:
        dx_km = (lon - olo) * km_per_deg_lon
        dy_km = (lat - ola) * km_per_deg_lat
        return w / 2 + dx_km * scale, h / 2 - dy_km * scale

    # grid
    draw.rectangle([0, 0, w - 1, h - 1], outline=(60, 90, 120))
    for g in range(0, w, 80):
        draw.line([(g, 0), (g, h)], fill=(28, 40, 58))
    for g in range(0, h, 80):
        draw.line([(0, g), (w, g)], fill=(28, 40, 58))

    if rmax > 0:
        outer = [tuple(int(round(v)) for v in geo_to_px(*p)) for p in range_ring_coords(ola, olo, rmax)]
        draw.polygon(outer, outline=(255, 120, 90), fill=(55, 38, 42))
        if rmin > 0.5:
            inner = [tuple(int(round(v)) for v in geo_to_px(*p)) for p in range_ring_coords(ola, olo, rmin)]
            draw.polygon(inner, outline=(255, 220, 120), fill=(30, 42, 58))
    cx, cy = geo_to_px(ola, olo)
    cx, cy = int(round(cx)), int(round(cy))
    r = 8
    draw.ellipse([cx - r, cy - r, cx + r, cy + r], fill=(0, 255, 200), outline=(200, 255, 240))

    title = f"{weapon.get('weapon_name', 'Weapon')} — range effects (schematic)"
    meta = f"Origin {ola:.4f}, {olo:.4f} | Max range {rmax} km | Inner ring ~{rmin:.1f} km"
    try:
        font = ImageFont.truetype("arial.ttf", 18)
        font_sm = ImageFont.truetype("arial.ttf", 13)
    except OSError:
        font = ImageFont.load_default()
        font_sm = font
    draw.text((margin, 12), title, fill=(220, 235, 255), font=font)
    draw.text((margin, 36), meta, fill=(160, 190, 220), font=font_sm)

    im.save(out_path, format="PNG")
    return True


def _esc(x: Any) -> str:
    if x is None:
        return ""
    return html.escape(str(x), quote=True)


def collect_image_paths(weapon: Dict[str, Any]) -> List[str]:
    paths: List[str] = []
    seen: set[str] = set()
    for img in weapon.get("images") or []:
        p = img.get("image_path")
        if p:
            key = p.lower()
            if key not in seen:
                seen.add(key)
                paths.append(p)
    pri = weapon.get("primary_image")
    if pri and pri.lower() not in seen:
        paths.insert(0, pri)
    return paths


INDIVIDUAL_SPEC_ROWS: Sequence[Tuple[str, str]] = tuple(SPEC_FIELDS)
AR_SPEC_ROWS: Sequence[Tuple[str, str]] = tuple((k, SPEC_FIELDS_AR.get(k, v)) for k, v in SPEC_FIELDS)


def _is_ar(lang_manager) -> bool:
    try:
        return bool(lang_manager and getattr(lang_manager, "is_arabic")())
    except Exception:
        return False


def _tr_data(lang_manager, value):
    try:
        return lang_manager.tr_data(value) if lang_manager else value
    except Exception:
        return value


def _resolve_spec_rows(meta: Optional[Dict[str, Any]], lang_manager=None) -> List[Tuple[str, str]]:
    requested = []
    if meta:
        requested = [str(k) for k in (meta.get("spec_fields") or []) if str(k).strip()]
    keys = set(requested) if requested else {k for k, _ in SPEC_FIELDS}
    rows = AR_SPEC_ROWS if _is_ar(lang_manager) else INDIVIDUAL_SPEC_ROWS
    return [(k, lbl) for k, lbl in rows if k in keys]

FULL_SPEC_COLUMNS: Sequence[Tuple[str, str]] = (
    ("weapon_name", "Weapon"),
    ("model", "Model"),
    ("category", "Category"),
    ("country", "Country"),
    ("status", "Status"),
    ("range_km", "Range km"),
    ("speed_mach", "Mach"),
    ("length_m", "Length m"),
    ("diameter_m", "Diam m"),
    ("weight_kg", "Weight kg"),
    ("warhead_type", "Warhead"),
    ("warhead_weight_kg", "WH kg"),
    ("guidance", "Guidance"),
    ("propulsion", "Propulsion"),
    ("platform", "Platform"),
    ("manufacturer", "Manufacturer"),
    ("intro_year", "Year"),
    ("unit_cost_usd", "Cost USD"),
    ("origin_lat", "Lat"),
    ("origin_lon", "Lon"),
    ("notes", "Notes"),
)


def _add_docx_specs_table(doc: Document, weapon: Dict[str, Any], lang_manager=None, meta: Optional[Dict[str, Any]] = None) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light List Accent 1"
    rows = _resolve_spec_rows(meta, lang_manager=lang_manager)
    for key, label in rows:
        row = table.add_row().cells
        row[0].text = label
        val = weapon.get(key)
        if key in {"category", "country", "status", "guidance", "propulsion", "warhead_type", "platform", "manufacturer"}:
            val = _tr_data(lang_manager, val)
        row[1].text = "" if val is None else str(val)
    doc.add_paragraph()


def _doc_closing_block(doc: Document, meta: Optional[Dict[str, Any]] = None, lang_manager=None) -> None:
    meta = meta or {}
    if not meta.get("include_closing", True):
        return
    closing = (meta.get("closing_notes") or "").strip()
    if not closing:
        return
    doc.add_heading("ملاحق / ملاحظات" if _is_ar(lang_manager) else "Appendix / notes", level=2)
    for line in closing.split("\n"):
        if line.strip():
            doc.add_paragraph(line.strip())


def _add_docx_images(doc: Document, weapon: Dict[str, Any], max_images: int = 12, lang_manager=None) -> None:
    doc.add_heading("معرض الصور" if _is_ar(lang_manager) else "Image gallery", level=2)
    paths = collect_image_paths(weapon)[:max_images]
    if not paths:
        doc.add_paragraph("لا توجد صور لهذا السلاح." if _is_ar(lang_manager) else "No images on file for this weapon.")
        return
    for p in paths:
        if Path(p).is_file():
            try:
                doc.add_picture(str(p), width=Inches(5.2))
                doc.add_paragraph()
            except Exception:
                doc.add_paragraph(f"(Could not embed image: {p})")


def write_individual_weapon_docx(
    weapon: Dict[str, Any],
    path: str,
    *,
    include_map: bool = False,
    meta: Optional[Dict[str, Any]] = None,
    lang_manager=None,
    map_png_path: Optional[str] = None,
) -> None:
    meta = meta or {}
    doc = Document()
    default_title = f"{'تقرير سلاح —' if _is_ar(lang_manager) else 'Weapon report —'} {weapon.get('weapon_name', 'Unknown')}"
    title = (meta.get("document_title") or "").strip() or default_title
    if meta.get("include_title", True):
        doc.add_heading(title, 0)
    intro = (meta.get("introduction") or "").strip()
    if intro and meta.get("include_intro", True):
        for line in intro.split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip())
    if _is_ar(lang_manager):
        doc.add_paragraph(
            f"الطراز: {weapon.get('model')} | الفئة: {_tr_data(lang_manager, weapon.get('category'))} | الدولة: {_tr_data(lang_manager, weapon.get('country'))}"
        )
        if meta.get("include_specs", True):
            doc.add_heading("المواصفات", level=2)
    else:
        doc.add_paragraph(
            f"Model: {weapon.get('model')} | Category: {weapon.get('category')} | Country: {weapon.get('country')}"
        )
        if meta.get("include_specs", True):
            doc.add_heading("Specifications", level=2)
    if meta.get("include_specs", True):
        _add_docx_specs_table(doc, weapon, lang_manager=lang_manager, meta=meta)
    if meta.get("include_gallery", True):
        _add_docx_images(doc, weapon, lang_manager=lang_manager)
    if include_map and meta.get("include_map", True):
        doc.add_heading("المدى والتأثيرات (خريطة)" if _is_ar(lang_manager) else "Range and effects (map diagram)", level=2)
        if map_png_path and Path(map_png_path).is_file():
            doc.add_picture(map_png_path, width=Inches(6.2))
        else:
            tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
            tmp.close()
            try:
                if build_range_effects_png(weapon, tmp.name):
                    doc.add_picture(tmp.name, width=Inches(6.2))
                else:
                    doc.add_paragraph("تعذر إنشاء صورة الخريطة." if _is_ar(lang_manager) else "Map diagram could not be generated.")
            finally:
                try:
                    os.unlink(tmp.name)
                except OSError:
                    pass
    _doc_closing_block(doc, meta, lang_manager=lang_manager)
    doc.save(path)


def write_multi_comparison_docx(
    weapons: Sequence[Dict[str, Any]],
    path: str,
    *,
    meta: Optional[Dict[str, Any]] = None,
    lang_manager=None,
) -> None:
    meta = meta or {}
    doc = Document()
    main = (meta.get("document_title") or "").strip() or ("تقرير مقارنة عدة أسلحة" if _is_ar(lang_manager) else "Multi-weapon comparison report")
    doc.add_heading(main, 0)
    intro = (meta.get("introduction") or "").strip()
    if intro:
        for line in intro.split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip())
    else:
        doc.add_paragraph("Side-by-side style sections with imagery and key parameters.")
    for idx, w in enumerate(weapons):
        doc.add_heading(str(w.get("weapon_name", "Weapon")), level=1)
        doc.add_paragraph(
            (
                f"الطراز: {w.get('model')} | {_tr_data(lang_manager, w.get('category'))} | {_tr_data(lang_manager, w.get('country'))}"
                if _is_ar(lang_manager)
                else f"Model: {w.get('model')} | {w.get('category')} | {w.get('country')}"
            )
        )
        tbl = doc.add_table(rows=1, cols=2)
        hdr = tbl.rows[0].cells
        hdr[0].text = "Field"
        hdr[1].text = "Value"
        for key, label in _resolve_spec_rows(meta, lang_manager=lang_manager):
            row = tbl.add_row().cells
            row[0].text = label
            val = w.get(key)
            if key in {"category", "country", "status", "guidance", "propulsion", "warhead_type", "platform", "manufacturer"}:
                val = _tr_data(lang_manager, val)
            row[1].text = str(val or "")
        _add_docx_images(doc, w, max_images=4, lang_manager=lang_manager)
        if idx < len(weapons) - 1:
            doc.add_page_break()
    _doc_closing_block(doc, meta, lang_manager=lang_manager)
    doc.save(path)


def write_category_grid_docx(
    weapons: Sequence[Dict[str, Any]], path: str, *, meta: Optional[Dict[str, Any]] = None, lang_manager=None
) -> None:
    meta = meta or {}
    by_cat: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for w in weapons:
        by_cat[str(w.get("category") or "Uncategorized")].append(w)
    doc = Document()
    main = (meta.get("document_title") or "").strip() or ("قائمة الأسلحة حسب الفئة" if _is_ar(lang_manager) else "Weapon inventory by category")
    doc.add_heading(main, 0)
    intro = (meta.get("introduction") or "").strip()
    if intro:
        for line in intro.split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip())
    for cat in sorted(by_cat.keys(), key=lambda s: s.lower()):
        doc.add_heading(str(_tr_data(lang_manager, cat)), level=1)
        rows = by_cat[cat]
        t = doc.add_table(rows=1, cols=5)
        h = t.rows[0].cells
        if _is_ar(lang_manager):
            h[0].text = "السلاح"
            h[1].text = "الطراز"
            h[2].text = "الدولة"
            h[3].text = "المدى (كم)"
            h[4].text = "الحالة"
        else:
            h[0].text = "Weapon"
            h[1].text = "Model"
            h[2].text = "Country"
            h[3].text = "Range km"
            h[4].text = "Status"
        for w in sorted(rows, key=lambda x: str(x.get("weapon_name", "")).lower()):
            r = t.add_row().cells
            r[0].text = str(w.get("weapon_name", ""))
            r[1].text = str(w.get("model", ""))
            r[2].text = str(_tr_data(lang_manager, w.get("country", "")))
            r[3].text = str(w.get("range_km", ""))
            r[4].text = str(_tr_data(lang_manager, w.get("status", "")))
        doc.add_paragraph()
    _doc_closing_block(doc, meta, lang_manager=lang_manager)
    doc.save(path)


def write_full_specs_docx(
    weapons: Sequence[Dict[str, Any]], path: str, *, meta: Optional[Dict[str, Any]] = None, lang_manager=None
) -> None:
    meta = meta or {}
    doc = Document()
    main = (meta.get("document_title") or "").strip() or ("قائمة المواصفات الكاملة للأسلحة" if _is_ar(lang_manager) else "Complete weapon specifications (all records)")
    doc.add_heading(main, 0)
    intro = (meta.get("introduction") or "").strip()
    if intro:
        for line in intro.split("\n"):
            if line.strip():
                doc.add_paragraph(line.strip())
    for w in weapons:
        doc.add_heading(str(w.get("weapon_name", "Weapon")), level=2)
        _add_docx_specs_table(doc, w, lang_manager=lang_manager, meta=meta)
        _add_docx_images(doc, w, max_images=6, lang_manager=lang_manager)
        doc.add_paragraph("—")
    _doc_closing_block(doc, meta, lang_manager=lang_manager)
    doc.save(path)


# --- PDF (ReportLab) --------------------------------------------------------


def _pdf_styles():
    styles = getSampleStyleSheet()
    return styles


def _pdf_append_closing(story: List[Any], styles, meta: Optional[Dict[str, Any]] = None) -> None:
    meta = meta or {}
    if not meta.get("include_closing", True):
        return
    closing = (meta.get("closing_notes") or "").strip()
    if not closing:
        return
    story.append(Spacer(1, 0.2 * inch))
    story.append(Paragraph("<b>Appendix / notes</b>", styles["Heading2"]))
    story.append(Paragraph(_esc(closing), styles["Normal"]))


def write_individual_weapon_pdf(
    weapon: Dict[str, Any],
    path: str,
    *,
    include_map: bool = False,
    meta: Optional[Dict[str, Any]] = None,
    lang_manager=None,
    map_png_path: Optional[str] = None,
) -> None:
    meta = meta or {}
    styles = _pdf_styles()
    story: List[Any] = []
    title = (meta.get("document_title") or "").strip() or weapon.get("weapon_name", "Weapon report")
    if meta.get("include_title", True):
        story.append(Paragraph(_esc(title), styles["Title"]))
    intro = (meta.get("introduction") or "").strip()
    if intro and meta.get("include_intro", True):
        for para in intro.split("\n\n"):
            if para.strip():
                story.append(Paragraph(_esc(para.strip()), styles["Normal"]))
    story.append(
        Paragraph(
            f"<b>Model:</b> {_esc(weapon.get('model'))} &nbsp;|&nbsp; "
            f"<b>Category:</b> {_esc(weapon.get('category'))} &nbsp;|&nbsp; "
            f"<b>Country:</b> {_esc(weapon.get('country'))}",
            styles["Normal"],
        )
    )
    story.append(Spacer(1, 0.15 * inch))
    if meta.get("include_specs", True):
        field_hdr = "الحقل" if _is_ar(lang_manager) else "Field"
        value_hdr = "القيمة" if _is_ar(lang_manager) else "Value"
        data = [[Paragraph(f"<b>{_esc(field_hdr)}</b>", styles["Normal"]), Paragraph(f"<b>{_esc(value_hdr)}</b>", styles["Normal"])]]
        rows = _resolve_spec_rows(meta, lang_manager=lang_manager)
        for key, label in rows:
            val = weapon.get(key)
            if key in {"category", "country", "status", "guidance", "propulsion", "warhead_type", "platform", "manufacturer"}:
                val = _tr_data(lang_manager, val)
            data.append([Paragraph(_esc(label), styles["Normal"]), Paragraph(_esc(val), styles["Normal"])])
        t = Table(data, colWidths=[2.0 * inch, 4.6 * inch])
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a3a52")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#334455")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.HexColor("#0f1826"), colors.HexColor("#121e2e")]),
                ]
            )
        )
        story.append(t)
        story.append(Spacer(1, 0.2 * inch))
    if meta.get("include_gallery", True):
        story.append(Paragraph("<b>المعرض</b>" if _is_ar(lang_manager) else "<b>Gallery</b>", styles["Heading2"]))
        imgs = collect_image_paths(weapon)[:10]
        if not imgs:
            story.append(Paragraph("No images on file.", styles["Normal"]))
        for p in imgs:
            if Path(p).is_file():
                try:
                    story.append(RLImage(str(p), width=4.8 * inch, height=3.2 * inch))
                    story.append(Spacer(1, 0.08 * inch))
                except Exception:
                    story.append(Paragraph(_esc(f"(Skipped image: {p})"), styles["Normal"]))
    if include_map and meta.get("include_map", True):
        story.append(PageBreak())
        story.append(Paragraph("<b>المدى والتأثيرات (خريطة)</b>" if _is_ar(lang_manager) else "<b>Range and effects (diagram)</b>", styles["Heading2"]))
        if map_png_path and Path(map_png_path).is_file():
            story.append(RLImage(map_png_path, width=6.5 * inch, height=4.2 * inch))
        else:
            tmp = tempfile.NamedTemporaryFile(suffix=".png", delete=False)
            tmp.close()
            try:
                if build_range_effects_png(weapon, tmp.name):
                    story.append(RLImage(tmp.name, width=6.5 * inch, height=4.2 * inch))
                else:
                    story.append(Paragraph("Diagram unavailable.", styles["Normal"]))
            finally:
                try:
                    os.unlink(tmp.name)
                except OSError:
                    pass
    _pdf_append_closing(story, styles, meta)
    doc = SimpleDocTemplate(path, pagesize=A4, title="Weapon report")
    doc.build(story)


def write_multi_comparison_pdf(
    weapons: Sequence[Dict[str, Any]], path: str, *, meta: Optional[Dict[str, Any]] = None, lang_manager=None
) -> None:
    meta = meta or {}
    styles = _pdf_styles()
    story: List[Any] = []
    main = (meta.get("document_title") or "").strip() or ("مقارنة عدة أسلحة" if _is_ar(lang_manager) else "Multi-weapon comparison")
    story.append(Paragraph(_esc(main), styles["Title"]))
    intro = (meta.get("introduction") or "").strip()
    if intro:
        for para in intro.split("\n\n"):
            if para.strip():
                story.append(Paragraph(_esc(para.strip()), styles["Normal"]))
    for i, w in enumerate(weapons):
        if i:
            story.append(PageBreak())
        story.append(Spacer(1, 0.08 * inch))
        story.append(Paragraph(_esc(w.get("weapon_name", "Weapon")), styles["Heading2"]))
        chips = []
        for key, label in _resolve_spec_rows(meta, lang_manager=lang_manager):
            val = w.get(key)
            if key in {"category", "country", "status", "guidance", "propulsion", "warhead_type", "platform", "manufacturer"}:
                val = _tr_data(lang_manager, val)
            chips.append(f"{_esc(label)}: {_esc(val)}")
            if len(chips) >= 4:
                break
        story.append(Paragraph(" | ".join(chips), styles["Normal"]))
        for p in collect_image_paths(w)[:3]:
            if Path(p).is_file():
                try:
                    story.append(RLImage(str(p), width=2.4 * inch, height=1.6 * inch))
                except Exception:
                    pass
    _pdf_append_closing(story, styles, meta)
    SimpleDocTemplate(path, pagesize=A4).build(story)


def write_category_grid_pdf(
    weapons: Sequence[Dict[str, Any]], path: str, *, meta: Optional[Dict[str, Any]] = None, lang_manager=None
) -> None:
    by_cat: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for w in weapons:
        by_cat[str(w.get("category") or "Uncategorized")].append(w)
    meta = meta or {}
    styles = _pdf_styles()
    story: List[Any] = []
    main = (meta.get("document_title") or "").strip() or ("قائمة الأسلحة حسب الفئة" if _is_ar(lang_manager) else "Weapons by category (grid list)")
    story.append(Paragraph(_esc(main), styles["Title"]))
    intro = (meta.get("introduction") or "").strip()
    if intro:
        for para in intro.split("\n\n"):
            if para.strip():
                story.append(Paragraph(_esc(para.strip()), styles["Normal"]))
    for cat in sorted(by_cat.keys(), key=lambda s: s.lower()):
        story.append(Paragraph(_esc(cat), styles["Heading2"]))
        rows: List[List[Any]] = [
            [
                Paragraph("<b>Weapon</b>", styles["Normal"]),
                Paragraph("<b>Model</b>", styles["Normal"]),
                Paragraph("<b>Country</b>", styles["Normal"]),
                Paragraph("<b>Range</b>", styles["Normal"]),
                Paragraph("<b>Status</b>", styles["Normal"]),
            ]
        ]
        for w in sorted(by_cat[cat], key=lambda x: str(x.get("weapon_name", "")).lower()):
            rows.append(
                [
                    Paragraph(_esc(w.get("weapon_name")), styles["Normal"]),
                    Paragraph(_esc(w.get("model")), styles["Normal"]),
                    Paragraph(_esc(w.get("country")), styles["Normal"]),
                    Paragraph(_esc(w.get("range_km")), styles["Normal"]),
                    Paragraph(_esc(w.get("status")), styles["Normal"]),
                ]
            )
        t = Table(rows, colWidths=[1.8 * inch, 1.5 * inch, 1.3 * inch, 0.9 * inch, 1.0 * inch])
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a3a52")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#334455")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ]
            )
        )
        story.append(t)
        story.append(Spacer(1, 0.2 * inch))
    _pdf_append_closing(story, styles, meta)
    SimpleDocTemplate(path, pagesize=landscape(A4)).build(story)


def write_full_specs_catalog_pdf(
    weapons: Sequence[Dict[str, Any]], path: str, *, meta: Optional[Dict[str, Any]] = None, lang_manager=None
) -> None:
    meta = meta or {}
    styles = _pdf_styles()
    keys = [k for k, _ in FULL_SPEC_COLUMNS]
    header = [Paragraph(f"<b>{_esc(lab)}</b>", styles["Normal"]) for _, lab in FULL_SPEC_COLUMNS]
    rows_pdf: List[List[Any]] = [header]
    for w in weapons:
        rows_pdf.append([Paragraph(_esc(w.get(k)), styles["Normal"]) for k in keys])
    cw = (10.8 * inch) / max(len(keys), 1)
    t = Table(rows_pdf, colWidths=[cw] * len(keys), repeatRows=1)
    t.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a3a52")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whitesmoke),
                ("GRID", (0, 0), (-1, -1), 0.2, colors.HexColor("#2b3f55")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 6),
            ]
        )
    )
    main = (meta.get("document_title") or "").strip() or ("قائمة المواصفات الكاملة للأسلحة" if _is_ar(lang_manager) else "Full specifications weapon list")
    story: List[Any] = [Paragraph(_esc(main), styles["Title"]), Spacer(1, 0.15 * inch)]
    intro = (meta.get("introduction") or "").strip()
    if intro:
        for para in intro.split("\n\n"):
            if para.strip():
                story.append(Paragraph(_esc(para.strip()), styles["Normal"]))
        story.append(Spacer(1, 0.1 * inch))
    story.append(t)
    _pdf_append_closing(story, styles, meta)
    SimpleDocTemplate(path, pagesize=landscape(A4)).build(story)
